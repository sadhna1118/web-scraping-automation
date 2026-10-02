"""
Data Export Module
Generates downloadable files in CSV, Excel (.xlsx), and JSON formats with professional formatting.
"""

import io
from typing import Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd


def export_to_csv(df: pd.DataFrame) -> bytes:
    """
    Export DataFrame to CSV bytes encoded with UTF-8 BOM ('utf-8-sig').
    Ensures seamless opening in Microsoft Excel without special character distortion.
    """
    if df is None or df.empty:
        return b""
    # index=False, utf-8-sig for Windows Excel compatibility
    csv_str = df.to_csv(index=False, encoding="utf-8-sig")
    return csv_str.encode("utf-8-sig")


def export_to_excel(df: pd.DataFrame, sheet_name: str = "Scraped_Data") -> bytes:
    """
    Export DataFrame to professional styled Excel (.xlsx) file in-memory using openpyxl.
    - Header row styled with dark navy fill, bold white text, and centered alignment.
    - Subtle cell borders.
    - Auto-adjusted column widths based on maximum text length.
    """
    if df is None or df.empty:
        return b""

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name[:31])  # Excel 31-char limit
        workbook = writer.book
        worksheet = writer.sheets[sheet_name[:31]]

        # Define Styles
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        # Format Header Row
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Format Data Rows & Zebra Striping
        for row_idx in range(2, len(df) + 2):
            is_even = (row_idx % 2 == 0)
            for col_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if is_even:
                    cell.fill = alt_fill

        # Auto-adjust column widths
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                # Consider first line if multiline
                val_str = val_str.split("\n")[0]
                if len(val_str) > max_len:
                    max_len = len(val_str)
            # Add padding and bound width
            adjusted_width = min(max(max_len + 4, 12), 60)
            worksheet.column_dimensions[col_letter].width = adjusted_width

    return output.getvalue()


def export_to_json(df: pd.DataFrame) -> bytes:
    """Export DataFrame to formatted JSON bytes."""
    if df is None or df.empty:
        return b"[]"
    json_str = df.to_json(orient="records", indent=2, force_ascii=False)
    return json_str.encode("utf-8")
